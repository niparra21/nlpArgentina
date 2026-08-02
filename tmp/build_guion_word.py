from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\nicop\Documents\GitHub\nlpArgentina")
SOURCE = ROOT / "docs" / "08-guion-exposicion-poster.md"
OUTPUT_DIR = ROOT / "outputs" / "guion-poster-20260802"
OUTPUT = OUTPUT_DIR / "guion-exposicion-poster.docx"

NAVY = "123B45"
TEAL = "0B6670"
PALE_BLUE = "EAF3F4"
PALE_TEAL = "E4F0F1"
PALE_GRAY = "F2F4F4"
INK = "202124"
MUTED = "5F6368"
WHITE = "FFFFFF"
GOLD = "8A5A00"
PALE_GOLD = "FFF4D6"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def paragraph_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def paragraph_left_border(paragraph, color=TEAL, size=18, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def define_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def add_definition(abstract_id, num_id, fmt, text, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.extend([tabs, ind])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        num.append(abs_id)
        numbering.append(num)

    add_definition(next_abs, next_num, "bullet", "•", "Symbol")
    bullet_num = next_num
    add_definition(next_abs + 1, next_num + 1, "decimal", "%1.")
    decimal_num = next_num + 1
    return bullet_num, decimal_num


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def clone_num_instance(doc: Document, source_num_id: int) -> int:
    """Create a fresh numbering instance so a later list restarts at 1."""
    numbering = doc.part.numbering_part.element
    source = numbering.find(f"{qn('w:num')}[@{qn('w:numId')}='{source_num_id}']")
    abstract_id = source.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    new_num_id = max(existing_ids or [0]) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), abstract_id)
    num.append(abs_id)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return new_num_id


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline_runs(paragraph, text: str, size=11, color=INK, default_bold=False):
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color, bold=default_bold)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(size - 0.5, 8), color=NAVY)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color, bold=default_bold)


def add_body_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    add_inline_runs(p, text)
    return p


def add_spoken_paragraph(doc, text):
    p = doc.add_paragraph(style="Spoken")
    paragraph_shading(p, PALE_BLUE)
    paragraph_left_border(p, TEAL, size=20, space=10)
    add_inline_runs(p, text, size=11.2, color=INK)
    return p


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, TEAL, 18, 10),
        "Heading 2": (13, TEAL, 14, 7),
        "Heading 3": (12, NAVY, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    spoken = styles.add_style("Spoken", 1)
    spoken.base_style = normal
    spoken.font.name = "Calibri"
    spoken.font.size = Pt(11.2)
    spoken.paragraph_format.left_indent = Inches(0.16)
    spoken.paragraph_format.right_indent = Inches(0.12)
    spoken.paragraph_format.space_before = Pt(4)
    spoken.paragraph_format.space_after = Pt(8)
    spoken.paragraph_format.line_spacing = 1.25
    spoken.paragraph_format.keep_together = False

    list_style = styles.add_style("Compact List", 1)
    list_style.base_style = normal
    list_style.font.name = "Calibri"
    list_style.font.size = Pt(11)
    list_style.paragraph_format.space_after = Pt(4)
    list_style.paragraph_format.line_spacing = 1.25
    list_style.paragraph_format.left_indent = Inches(0.30)
    list_style.paragraph_format.first_line_indent = Inches(-0.22)

    note_style = styles.add_style("Note", 1)
    note_style.base_style = normal
    note_style.font.name = "Calibri"
    note_style.font.size = Pt(10.5)
    note_style.font.color.rgb = rgb(NAVY)
    note_style.paragraph_format.left_indent = Inches(0.12)
    note_style.paragraph_format.right_indent = Inches(0.12)
    note_style.paragraph_format.space_before = Pt(6)
    note_style.paragraph_format.space_after = Pt(8)
    note_style.paragraph_format.line_spacing = 1.2


def configure_page(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "Guion de exposición  |  Segunda Escuela Sudamericana de NLP 2026"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    for run in hp.runs:
        set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("Nicole Tatiana Parra Valverde   |   Página ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(fp)


def add_title_block(doc: Document):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("GUION DE EXPOSICIÓN")
    set_run_font(run, size=10.5, color=TEAL, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("Diseño arquitectónico mínimo para sistemas RAG observables en español latinoamericano")
    set_run_font(run, size=24, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    run = subtitle.add_run("Nicole Tatiana Parra Valverde  |  Guía para presentación y conversación")
    set_run_font(run, size=12.5, color=MUTED)

    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120], indent_dxa=120)
    remove_table_borders(table)
    metrics = [
        ("3–4 min", "recorrido principal"),
        ("30 s", "versión breve"),
        ("2 bloques", "sesión del martes"),
    ]
    for cell, (value, label) in zip(table.rows[0].cells, metrics):
        set_cell_shading(cell, PALE_TEAL)
        set_cell_margins(cell, top=130, bottom=130, start=160, end=160)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(value)
        set_run_font(r, size=15, color=TEAL, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(label)
        set_run_font(r2, size=9, color=MUTED)

    p = doc.add_paragraph(style="Note")
    paragraph_shading(p, PALE_BLUE)
    paragraph_left_border(p, TEAL, size=20, space=10)
    lead = p.add_run("Idea central. ")
    set_run_font(lead, size=10.5, color=NAVY, bold=True)
    rest = p.add_run(
        "La observabilidad permite explicar por qué una respuesta RAG falla, no solamente medir si el resultado final parece correcto."
    )
    set_run_font(rest, size=10.5, color=NAVY)


def parse_markdown_into_doc(doc: Document, text: str, bullet_num: int, decimal_num: int):
    lines = text.splitlines()
    i = 0
    skipped_title = False
    current_decimal_num = decimal_num
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue

        if line.startswith("# ") and not skipped_title:
            skipped_title = True
            i += 1
            continue

        if line.startswith("## "):
            title = line[3:].strip()
            current_decimal_num = clone_num_instance(doc, decimal_num)
            p = doc.add_paragraph(style="Heading 1")
            add_inline_runs(p, title, size=16, color=TEAL, default_bold=True)
            if title == "Verificación pendiente del número de orden":
                paragraph_shading(p, PALE_GOLD)
            i += 1
            continue

        if line.startswith("### "):
            title = line[4:].strip()
            current_decimal_num = clone_num_instance(doc, decimal_num)
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, title, size=13, color=TEAL, default_bold=True)
            i += 1
            continue

        if line.startswith(">"):
            parts = []
            while i < len(lines):
                candidate = lines[i].strip()
                if not candidate.startswith(">"):
                    break
                parts.append(candidate[1:].strip())
                i += 1
            add_spoken_paragraph(doc, " ".join(parts))
            continue

        if re.match(r"^-\s+", line):
            text_value = re.sub(r"^-\s+", "", line)
            p = doc.add_paragraph(style="Compact List")
            add_inline_runs(p, f"•  {text_value}")
            i += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            match = re.match(r"^(\d+)\.\s+(.*)", line)
            p = doc.add_paragraph(style="Compact List")
            add_inline_runs(p, f"{match.group(1)}.  {match.group(2)}")
            i += 1
            continue

        parts = [line]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate or candidate.startswith(("#", ">", "- ")) or re.match(r"^\d+\.\s+", candidate):
                break
            parts.append(candidate)
            i += 1
        p = add_body_paragraph(doc, " ".join(parts))
        if "El correo de aceptación" in p.text:
            paragraph_shading(p, PALE_GOLD)
            paragraph_left_border(p, GOLD, size=20, space=10)
            p.style = doc.styles["Note"]


def append_sources(doc: Document):
    p = doc.add_paragraph(style="Heading 1")
    add_inline_runs(p, "Fuentes de organización", size=16, color=TEAL, default_bold=True)

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    add_hyperlink(p1, "Programa y sesiones de pósters de la Escuela", "https://south-american-nlp-school.dc.uba.ar/")

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    add_hyperlink(p2, "Anuncio institucional de ECI", "https://eci.dc.uba.ar/segunda-escuela-sudamericana-de-nlp/")

    p3 = doc.add_paragraph(style="Note")
    p3.add_run("Contacto de la organización: ")
    add_hyperlink(p3, "escuelalatamnlp@googlegroups.com", "mailto:escuelalatamnlp@googlegroups.com")


def preset_audit(doc: Document):
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert doc.styles["Normal"].font.name == "Calibri"
    assert round(doc.styles["Normal"].font.size.pt, 1) == 11.0
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 1 and len(table.columns) == 3
    assert all("Heading" not in p.style.name or p.text.strip() for p in doc.paragraphs)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    bullet_num, decimal_num = define_numbering(doc)
    add_title_block(doc)

    md = SOURCE.read_text(encoding="utf-8")
    parse_markdown_into_doc(doc, md, bullet_num, decimal_num)
    append_sources(doc)
    preset_audit(doc)

    props = doc.core_properties
    props.title = "Guion de exposición del póster RAG observable"
    props.subject = "Segunda Escuela Sudamericana de NLP 2026"
    props.author = "Nicole Tatiana Parra Valverde"
    props.keywords = "RAG, observabilidad, NLP, póster, guion"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
